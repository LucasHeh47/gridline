import datetime, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

out = sys.argv[1]
wb = Workbook()
ws = wb.active
ws.title = "Sales Data"
hdr = ["Date", "Region", "Rep", "Units", "Unit Price", "Revenue", "Tax"]
ws.append(hdr)
regions = ["North", "South", "East", "West"]
reps = ["Ann", "Bob", "Cara", "Dev"]
for i in range(20):
    d = datetime.date(2025, 1, 1) + datetime.timedelta(days=i * 7)
    ws.append([d, regions[i % 4], reps[(i * 3) % 4], 10 + (i * 7) % 50, [4.5, 12, 30, 7.25][i % 4], None, None])
    r = i + 2
    ws.cell(row=r, column=6, value=f"=D{r}*E{r}")
    ws.cell(row=r, column=7, value=f"=ROUND(F{r}*Rate,2)")
ws["I1"] = "Tax rate"
ws["J1"] = 0.0825
ws["J1"].number_format = "0.00%"
wb.defined_names["Rate"] = DefinedName("Rate", attr_text="'Sales Data'!$J$1")
ws["I3"] = "Total revenue"
ws["J3"] = "=SUM(F2:F21)"
ws["I4"] = "North revenue"
ws["J4"] = '=SUMIFS(F2:F21,B2:B21,"North")'
ws["I5"] = "Lookup Cara"
ws["J5"] = '=IFERROR(VLOOKUP("Cara",C2:F21,4,FALSE),"none")'
ws["I6"] = "Avg units"
ws["J6"] = "=AVERAGE(D2:D21)"
for c in "JKL":
    pass
for col in range(1, 8):
    cell = ws.cell(row=1, column=col)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="305496")
    cell.alignment = Alignment(horizontal="center")
    cell.border = Border(bottom=Side(style="medium", color="000000"))
for r in range(2, 22):
    ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"
    ws.cell(row=r, column=5).number_format = '"$"#,##0.00'
    ws.cell(row=r, column=6).number_format = '"$"#,##0.00'
    ws.cell(row=r, column=7).number_format = '"$"#,##0.00'
ws.column_dimensions["A"].width = 14
ws.column_dimensions["C"].width = 10
ws.column_dimensions["I"].width = 18
ws.row_dimensions[1].height = 24
ws.freeze_panes = "A2"
ws.auto_filter.ref = "A1:G21"
ws.conditional_formatting.add("D2:D21", CellIsRule(operator="greaterThan", formula=["40"], fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(color="9C0006")))
ws.conditional_formatting.add("F2:F21", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
ws.conditional_formatting.add("E2:E21", DataBarRule(start_type="min", end_type="max", color="638EC6"))
ws.conditional_formatting.add("B2:B21", FormulaRule(formula=['$B2="West"'], font=Font(italic=True, color="0070C0")))
ws["C2"].comment = Comment("Top performer", "Manager")
ws["I8"] = "Docs"
ws["I8"].hyperlink = "https://example.com/report"
dv = DataValidation(type="list", formula1='"North,South,East,West"')
ws.add_data_validation(dv)
dv.add("B2:B21")

chart = BarChart()
chart.title = "Revenue by row"
chart.y_axis.title = "Revenue"
data = Reference(ws, min_col=6, min_row=1, max_row=21)
cats = Reference(ws, min_col=1, min_row=2, max_row=21)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
ws.add_chart(chart, "L2")

s2 = wb.create_sheet("Summary")
s2["A1"] = "Region"
s2["B1"] = "Revenue"
for i, reg in enumerate(regions):
    s2.cell(row=i + 2, column=1, value=reg)
    s2.cell(row=i + 2, column=2, value=f"=SUMIFS('Sales Data'!F:F,'Sales Data'!B:B,A{i+2})")
s2["A7"] = "Merged title"
s2.merge_cells("A7:C8")
s2["A7"].alignment = Alignment(horizontal="center", vertical="center")
s2["A7"].font = Font(size=16, bold=True)
pie = PieChart()
pie.title = "Share"
pie.add_data(Reference(s2, min_col=2, min_row=1, max_row=5), titles_from_data=True)
pie.set_categories(Reference(s2, min_col=1, min_row=2, max_row=5))
s2.add_chart(pie, "D2")
line = LineChart()
line.title = "Trend"
line.add_data(Reference(s2, min_col=2, min_row=1, max_row=5), titles_from_data=True)
s2.add_chart(line, "D20")
wb.save(out)
print("wrote", out)
